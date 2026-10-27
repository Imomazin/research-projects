"""Assemble the analysis workbook from the analysis CSV files (no values typed by hand)."""
import json
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
AN = HERE / "Results" / "Analysis"
OUT = HERE.parents[1] / "Manuscript" / "Privacy Trust and Intelligent Systems Analysis.xlsx"
HEAD = PatternFill("solid", fgColor="E8EEF4")
BOLD = Font(name="Calibri", bold=True, size=10)
BODY = Font(name="Calibri", size=10)

DICTIONARY = [
    ("dataset", "Task identifier: maintenance (AI4I), semiconductor (SECOM), steel, power, gas, iot (RT-IoT2022), bosch, hai"),
    ("design", "Evaluation design: simulated clients, machines, time holdout, machine holdout, sessions"),
    ("seed", "Confirmatory seed; seed 0 was used only for tuning"),
    ("model", "reference, linear, central (pooled network), local, fedavg, fedprox, scaffold, trimmed, median, dp eps N"),
    ("method", "Certification rule: default, nonprivate_plugin, nonprivate_ltt, plugin (private), ucb_without_noise (private), proposed (PARC)"),
    ("test_set", "test (in-distribution held-out rows) or a shifted set: year 2014, year 2015, later periods, unseen machine Mxx, later sessions"),
    ("metric", "pr_auc, roc_auc, balanced_accuracy, macro_f1, mcc, recall_rare, brier, nll, ece, rmse (physical units), mae, r2, miss_rate, false_alarm_rate, escalation_rate, coverage, exceedance, half_width, joint_error"),
    ("value", "Metric value for one run; proportions are on a 0 to 1 scale"),
    ("mean, ci_low, ci_high", "Mean over seeds and percentile bootstrap 95% interval (10,000 resamples); intervals describe seed variation, not site variation"),
    ("epsilon_*", "Record-level privacy budgets at delta 1e-5 (sequential values at 2e-5)"),
]


def write_df(ws, df, start_row=1):
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(row=start_row, column=j, value=str(c))
        cell.font, cell.fill = BOLD, HEAD
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            ws.cell(row=i, column=j, value=v).font = BODY
    for j, c in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = min(48, max(10, len(str(c)) + 2))
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return start_row + len(df) + 2


def section(ws, row, title, df):
    ws.cell(row=row, column=1, value=title).font = Font(name="Calibri", bold=True, size=11)
    return write_df(ws, df, row + 1)


def read(name):
    p = AN / f"{name}.csv"
    return pd.read_csv(p) if p.exists() and p.stat().st_size > 0 else pd.DataFrame()


def calibration_evidence():
    rows = []
    for p in sorted((HERE / "Results" / "Main").glob("*/*.json")):
        d = json.loads(p.read_text())
        if d.get("kind") == "binary":
            pos = sum(c["calibration_positive"] for c in d["clients"])
            rows.append({"dataset": d["dataset"], "design": d["design"], "seed": d["seed"], "clients": len(d["clients"]),
                         "calibration_positives": pos,
                         "calibration_negatives": sum(c["calibration"] for c in d["clients"]) - pos})
    return pd.DataFrame(rows)


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "About"
    meta = [("Title", "Privacy, Trust & Intelligent Systems in Industry 6.0: analysis workbook"),
            ("Author", "Imo Enang"), ("Built (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")),
            ("Source", "Results/Analysis CSV files produced by Analyse Results.py from per-run JSON outputs"),
            ("Repository", "Imomazin/research-projects, branch research/industrial-privacy-trust"),
            ("Frozen protocol commit", "1dc2dbf")]
    for i, (k, v) in enumerate(meta, 1):
        ws.cell(row=i, column=1, value=k).font = BOLD
        ws.cell(row=i, column=2, value=v).font = BODY
    r = len(meta) + 2
    ws.cell(row=r, column=1, value="Data dictionary").font = Font(name="Calibri", bold=True, size=11)
    for i, (k, v) in enumerate(DICTIONARY, r + 1):
        ws.cell(row=i, column=1, value=k).font = BOLD
        ws.cell(row=i, column=2, value=v).font = BODY
        ws.cell(row=i, column=2).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110

    ws = wb.create_sheet("Run Results")
    u = read("Utility")
    u = u[u.client.isna()].drop(columns=["client"]) if len(u) else u
    row = section(ws, 1, "Utility per run (global and shifted tests)", u)

    ws = wb.create_sheet("Decision Runs")
    section(ws, 1, "Certification outcomes per run", read("Risk Control"))

    ws = wb.create_sheet("Summaries")
    row = section(ws, 1, "Utility summary by task, model and test set", read("Utility Summary"))
    rc = read("Risk Control Summary")
    row = section(ws, row, "Decision summary by certification rule", rc)
    row = section(ws, row, "Membership inference summary", read("Membership Summary"))
    row = section(ws, row, "HAI event detection summary", read("HAI Events Summary"))
    row = section(ws, row, "Gas turbine recalibration on 2014, tested on 2015", read("Gas Recalibration Summary"))

    ws = wb.create_sheet("Tests")
    hyp = json.loads((AN / "Hypotheses.json").read_text())
    rows = []
    for k, v in hyp.items():
        if isinstance(v, dict):
            for kk, vv in v.items():
                if not isinstance(vv, (list, dict)):
                    rows.append({"hypothesis": k, "quantity": kk, "value": vv})
    row = section(ws, 1, "Pre-specified hypotheses H2 to H4 with Holm adjustment", pd.DataFrame(rows))
    row = section(ws, row, "Shifted minus in-distribution realised risk (H4 pairs)", read("Shift Differences"))
    row = section(ws, row, "Paired utility comparisons against FedAvg (Wilcoxon, Benjamini-Hochberg)",
                  read("Paired Utility Comparisons"))
    mc = HERE / "Results" / "Guarantee Validation" / "Monte Carlo Validation.json"
    if mc.exists():
        row = section(ws, row, "Monte Carlo validation of the guarantee (H1)", pd.DataFrame(json.loads(mc.read_text())["rows"]))

    ws = wb.create_sheet("Privacy and Costs")
    row = section(ws, 1, "Privacy accounting summary", read("Privacy Summary"))
    row = section(ws, row, "Evidence floor (exact computation)", read("Evidence Floor"))
    row = section(ws, row, "Computation and communication", read("Cost Summary"))

    ws = wb.create_sheet("Robustness")
    row = 1
    for name, title in [("Sensitivity Summary", "Sensitivity and attack experiments"),
                        ("Calibration Sweep", "Calibration-only privacy sweep (added secondary analysis)"),
                        ("Reconstruction", "Analytic reconstruction from single-record updates")]:
        df = read(name)
        if len(df):
            row = section(ws, row, title, df)

    ws = wb.create_sheet("Figure Data")
    src = HERE.parent.parent / "Manuscript" / "Manuscript Source.json"
    where = {
        "Research Framework": "Conceptual diagram; no data",
        "Architecture and Protocol": "Conceptual diagram; no data",
        "Privacy Utility": "Summaries: utility summary, models FedAvg and dp eps 1/2/4/8, test set test",
        "Model Comparison": "Summaries: utility summary, all models, test set test",
        "Ablation of Certification Rules": "Decision Runs: model fedavg, kind binary, metrics miss_rate, false_alarm_rate, escalation_rate",
        "Shift and Unseen Machines": "Decision Runs: method proposed, model fedavg, shifted test sets; Tests: shift differences",
        "Robustness Under Attack and Corruption": "Robustness: sensitivity and attack experiments",
        "Guarantee Validity and Evidence Floor": "Tests: Monte Carlo validation; Privacy and Costs: evidence floor; Robustness: calibration-only sweep",
    }
    tables = {1: "Literature evidence matrix in the repository (Literature folder)", 2: "Source register in the repository (Datasets folder)",
              3: "Study configuration in the repository", 4: "Summaries: utility summary", 5: "Summaries: decision summary",
              6: "Privacy and Costs: evidence floor", 7: "Tests: pre-specified hypotheses", 8: "Robustness: sensitivity and attack experiments",
              9: "Privacy and Costs: privacy accounting summary", 10: "Privacy and Costs: computation and communication",
              11: "Derived from Tables 4 to 10", 12: "Summaries: utility summary", 13: "Tuned hyperparameters in the repository"}
    rows = []
    if src.exists():
        import re
        for b in json.loads(src.read_text())["blocks"]:
            if b["type"] == "figure":
                cap = b.get("caption", "")
                m = re.match(r"\*\*(Figure \d+)\.\*\*\s*(.*)", cap)
                name = Path(b.get("src", b.get("path", ""))).stem
                rows.append({"Item": m.group(1) if m else name, "Caption": (m.group(2) if m else cap)[:250],
                             "Image file": name, "Source data in this workbook": where.get(name, "")})
            if b["type"] == "table":
                cap = b.get("caption", "")
                m = re.match(r"\*\*Table (\d+)\.\*\*\s*(.*)", cap)
                if m:
                    rows.append({"Item": f"Table {m.group(1)}", "Caption": m.group(2)[:250], "Image file": "",
                                 "Source data in this workbook": tables.get(int(m.group(1)), "")})
    row = section(ws, 1, "Where each manuscript figure and table draws its data", pd.DataFrame(rows))
    row = section(ws, row, "Calibration positives and negatives per run (evidence available for certification)", calibration_evidence())
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
