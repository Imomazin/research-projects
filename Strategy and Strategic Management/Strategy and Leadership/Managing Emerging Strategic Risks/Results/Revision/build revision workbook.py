"""Write the revision workbook in a grey and white layout with a filter on every table.

Every table is read from the revision outputs, the validation files, the reproduction files or
the original results; nothing is typed by hand except the read-me text, the reproducibility
record and the column labels. Output: Results/Managing Emerging Strategic Risks Revision Workbook.xlsx
"""
import json
import re
import subprocess
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
RES = HERE.parent
PROJ = RES.parent
VAL, REP, DER = PROJ / "Validation", PROJ / "Reproduction", PROJ / "Datasets" / "Derived"
OUT = RES / "Managing Emerging Strategic Risks Revision Workbook.xlsx"
N = json.loads((PROJ / "Manuscript" / "Revision Numbers.json").read_text())
R = json.loads((HERE / "Revision Results.json").read_text())


def git(*a):
    return subprocess.run(["git", *a], cwd=PROJ, capture_output=True, text=True).stdout.strip()


commit = git("rev-parse", "--short", "HEAD")
branch = git("rev-parse", "--abbrev-ref", "HEAD")
rep = N.get("reproduction", {})
csv = lambda p, **k: pd.read_csv(p, low_memory=False, **k)

ABBR = {"cik": "CIK", "sic2": "SIC2", "sic": "SIC", "p": "P value", "q bh": "BH q value", "p holm": "Holm p value", "p wild": "Wild bootstrap p value",
        "p t": "P value, t reference", "p hc1": "P value, HC1", "coef": "Coefficient", "se": "Standard error", "n": "N", "r2": "R squared",
        "ci low": "CI low", "ci high": "CI high", "ci low t": "CI low, t reference", "ci high t": "CI high, t reference", "sd y": "SD of outcome",
        "fy": "FY", "lm": "Language model", "cv": "Coefficient of variation", "id": "ID", "abs difference": "Absolute difference"}


def label(c):
    s = re.sub(r"_+", " ", str(c)).strip()
    if s.lower() in ABBR:
        return ABBR[s.lower()]
    return s[:1].upper() + s[1:]


def tidy(df):
    df = df.copy()
    df.columns = [label(c) for c in df.columns]
    seen = {}
    cols = []
    for c in df.columns:                      # Excel tables need unique header text
        seen[c] = seen.get(c, 0) + 1
        cols.append(c if seen[c] == 1 else f"{c} {seen[c]}")
    df.columns = cols
    for c in df.columns:
        if df[c].dtype == object:
            df[c] = df[c].map(lambda v: json.dumps(v) if isinstance(v, (dict, list)) else v)
    return df


def flat(d, prefix=""):
    for k, v in d.items():
        if isinstance(v, dict):
            yield from flat(v, f"{prefix}{k} / ")
        else:
            yield f"{prefix}{k}", v


S = {}
# ------------------------------------------------------------------ 1. read me and reproducibility record
readme = pd.DataFrame([
    ["Title", "When every risk is material: disclosed risk convergence, documented response and organisational resilience in US firms, 2017 to 2025 (revision)"],
    ["Author", "Imo J. Enang"],
    ["Purpose", "Source tables for the revised manuscript, the reproduction of the first version, the validation of the response measures and the corrections register"],
    ["Sample", "787 US non-financial firms; 6,214 10-K filings, fiscal 2017 to 2025"],
    ["Numbers", "Every value comes from versioned scripts in the repository; none is typed by hand"],
    ["Human coding", "Reference labels in the validation sheets were assigned by an AI assistant applying the coding manual. They are not independent human annotation. Blank sheets for two human coders are in Validation/"],
    ["Sheet: Data Dictionary", "Variable definitions (first version and revision) and the dataset inventory with checksums"],
    ["Sheet: Reconciliation", "Reproduction of every quantity in the first manuscript from raw data, file comparison and the independent re-implementation"],
    ["Sheet: Corrections Register", "Issue, original finding or method, verification result, correction and effect on conclusions"],
    ["Sheet: Sample and Validation", "Sample selection, descriptives and validation of four documented-response measures (Tables 2 and 3)"],
    ["Sheet: Hypothesis Tests", "Principal models with full inference, response-measure comparisons, marginal effects, placebo and category models (Tables 4 and 5, Figure 5)"],
    ["Sheet: Robustness", "Recovery and interaction scrutiny, robustness grids, attrition checks and alternative convergence measures (Tables 5, A3, A4)"],
    ["Sheet: Convergence and Shocks", "Dispersion under alternative counts, determinants of breadth, shock-channel models and equivalence tests (Table 6, Figure 3)"],
    ["Sheet: Figure Data", "Data behind Figures 2 to 5"],
], columns=["Item", "Description"])
record = pd.DataFrame([
    ["Repository", "Imomazin/research-projects (GitHub); folder Strategy and Strategic Management/Strategy and Leadership/Managing Emerging Strategic Risks"],
    ["Revision branch", f"{branch}; commit {commit} when this workbook was built (the final commit is recorded in the repository README)"],
    ["First-version branch and commit", "research/strategic-risk at ed1a1df; data-processing commit fe19a29 is on that branch"],
    ["Raw data branch", "data/strategic-risk (filing text in nine parts, XBRL company facts, submissions)"],
    ["Raw data versions", "Retrieval manifests with SHA-256 checksums in Datasets/Raw Macro/ and on the data branch; macro files re-downloaded on 25 October 2026 and byte-identical"],
    ["Language model", "MoritzLaurer/DeBERTa-v3-base-mnli-fever-anli at revision 6f5cf0a2b59cabb106aca4c287eed12e357e90eb; threshold 0.5; run in GitHub Actions stage corpusnli"],
    ["Software", f"Python {rep.get('python', sys.version.split()[0])}; pandas {rep.get('pandas', pd.__version__)}; numpy {rep.get('numpy', '')}; statsmodels {rep.get('statsmodels', '')}; linearmodels {rep.get('linearmodels', '')}; scikit-learn {rep.get('sklearn', '')}; transformers and torch in CI; Node docx for the manuscript"],
    ["Run order", "1 Reproduction/reproduce study.py --data <data branch> --work <folder>; 2 Analysis/build response measures.py and build convergence measures.py; "
                  "3 Analysis/export response candidates.py then CI stage corpusnli; 4 Validation/validate response measures.py; 5 Analysis/revision models.py; "
                  "6 Analysis/make revision figures.py; 7 Manuscript/collect revision numbers.py; 8 Reproduction/build corrections register.py; "
                  "9 Manuscript/build blocks revision.py and node build manuscript revision.js; 10 this workbook"],
    ["Random seeds", "Sample draw seed fixed in Data Collection/draw sample.py; validation sample seed 20261025; wild bootstrap seed 0"],
    ["Network access", "SEC, BLS, FRED, Crossref, OpenAlex and Hugging Face are reached only from GitHub Actions; the analysis runs offline from the committed files"],
    ["Limitations", "Human double-coding outstanding; zero-shot model untuned; nominal accounting outcomes; 53 uneven industry clusters; associations only"],
], columns=["Item", "Record"])
S["Read Me"] = [("Workbook contents", readme), ("Reproducibility record", record)]

# ------------------------------------------------------------------ 2. data dictionary and inventory
inv = PROJ / "Datasets" / "Revision Dataset Manifest.csv"
S["Data Dictionary"] = [("Revision variables", csv(PROJ / "Datasets" / "Revision Data Dictionary.csv")),
                        ("First-version variables", csv(PROJ / "Datasets" / "Data Dictionary.csv"))] + \
                       ([("Dataset inventory (revision files)", csv(inv))] if inv.exists() else []) + \
                       [("Dataset inventory (first version)", csv(PROJ / "Datasets" / "Dataset Manifest.csv"))]

# ------------------------------------------------------------------ 3. reconciliation
rec = []
if (REP / "Reconciliation Report.csv").exists():
    rec.append(("Reconciliation of every quantity in the first manuscript: original value, reproduced value, explanation and correction",
                csv(REP / "Reconciliation Report.csv")))
if (REP / "File Comparison.csv").exists():
    rec.append(("File-level comparison of reproduced outputs with the committed outputs at ed1a1df", csv(REP / "File Comparison.csv")))
if (REP / "Reproduction Run Log.csv").exists():
    rec.append(("Reproduction run log", csv(REP / "Reproduction Run Log.csv")))
rec += [("Independent re-implementation written without the analysis code", csv(REP / "Independent Checks.csv")),
        ("Exit and attrition check", csv(REP / "Exit and Attrition Check.csv"))]
S["Reconciliation"] = rec

# ------------------------------------------------------------------ 4. corrections register
S["Corrections Register"] = [("Corrections register", csv(REP / "Corrections Register.csv"))]

# ------------------------------------------------------------------ 5. sample and validation
S["Sample and Validation"] = [
    ("Sample selection log", csv(PROJ / "Datasets" / "Sampling Frame" / "Sample Selection Log.csv")),
    ("Pandemic episode descriptive statistics (winsorised)", csv(RES / "Descriptives Episode A.csv").rename(columns={"Unnamed: 0": "Variable"})),
    ("2022 episode descriptive statistics (winsorised)", csv(RES / "Descriptives Episode B.csv").rename(columns={"Unnamed: 0": "Variable"})),
    ("Share of firms with any documented response by rule and fiscal year (per cent)",
     pd.DataFrame(N["any_response_pct"]).rename_axis("Fiscal year").reset_index()),
    ("Validation statistics by rule (weighted figures are population estimates; labels from AI reference coding)", csv(VAL / "Response Validation Statistics.csv")),
    ("Validation by category", csv(VAL / "Response Validation by Category.csv")),
    ("Error analysis", csv(VAL / "Response Validation Errors.csv")),
    ("Confusion counts", csv(VAL / "Response Validation Confusion.csv")),
    ("Population of sentence-category pairs by stratum", csv(VAL / "Response Validation Universe Counts.csv")),
    ("Validation sample with reference coding (AI-assigned, not human)",
     csv(VAL / "Response Validation Sample.csv").merge(csv(VAL / "Response Validation Reference Coding.csv").drop(columns=["sentence", "cik", "fiscal_year", "category", "division"], errors="ignore"), on="item_id")),
]

# ------------------------------------------------------------------ 6. hypothesis tests
hyp = [("Principal models: CR1 by two-digit SIC, t reference, HC1, wild bootstrap, Holm and BH", csv(HERE / "Principal Models Full Inference.csv")),
       ("Response-measure comparison for H2 and H3 (BH across the family)", csv(HERE / "Response Measure Variants.csv")),
       ("Marginal effect of breadth at levels of coverage, by response measure", csv(HERE / "Marginal Effects by Response Measure.csv"))]
if (HERE / "Placebo Episode.csv").exists():
    hyp.append(("Placebo episode: fiscal 2017 predictors, fiscal 2018 to 2019 outcomes", csv(HERE / "Placebo Episode.csv")))
hyp += [("Category-level response models", csv(HERE / "Category Response Models.csv")),
        ("First-version episode models (reproduced)", csv(RES / "Episode Models.csv"))]
S["Hypothesis Tests"] = hyp

# ------------------------------------------------------------------ 7. robustness
S["Robustness"] = [
    ("Recovery association: alternative explanations", csv(HERE / "Recovery Investigation.csv")),
    ("Breadth-by-coverage interaction: scrutiny", csv(HERE / "Interaction Investigation.csv")),
    ("Robustness grid, breadth alone (revision)", csv(HERE / "Robustness Grid Breadth Alone.csv")),
    ("Robustness grid (first version, reproduced)", csv(RES / "Robustness Grid.csv")),
    ("Attrition and added pre-registered checks", csv(HERE / "Preregistered Checks Added.csv")),
    ("Survival to fiscal 2022: logit coefficients and p-values", pd.DataFrame({"Term": list(R["survival_logit"]["coef"]),
                                                                            "Coefficient": list(R["survival_logit"]["coef"].values()),
                                                                            "P value": [R["survival_logit"]["p"][k] for k in R["survival_logit"]["coef"]]})),
    ("Alternative convergence measures in the H1 model", csv(HERE / "Alternative Convergence Measures.csv")),
    ("Cluster sizes, pandemic episode", csv(HERE / "Cluster Sizes Episode A.csv")),
]

# ------------------------------------------------------------------ 8. convergence and shocks
sat = pd.DataFrame(R["saturation"]).drop(columns=["below_ceiling_firms"], errors="ignore").rename_axis("Fiscal year").reset_index() \
    if isinstance(R["saturation"].get("cv_breadth"), dict) else pd.DataFrame()
comp = pd.DataFrame([{**{k: v for k, v in c.items() if not isinstance(v, list)},
                      "ci95 low": c["ci95"][0], "ci95 high": c["ci95"][1], "ci90 low": c["ci90"][0], "ci90 high": c["ci90"][1]} for c in R["compounding_equivalence"]])
S["Convergence and Shocks"] = [
    ("Measures by fiscal year: mean, SD, coefficient of variation, P90 minus P10", csv(HERE / "Measure Trends and Dispersion.csv")),
    ("Saturation indicators by fiscal year", sat),
    ("Dispersion among firms that never reached nine categories", csv(HERE / "Saturation Below Ceiling Firms.csv")),
    ("Determinants of breadth (R squared by predictor set)", csv(HERE / "Determinants of Breadth.csv")),
    ("Exposure-by-shock models: timings, fixed effects and clustering", csv(HERE / "Shock Models Revision.csv")),
    ("Compounding: equivalence tests against the smallest significant single-channel effect", comp),
    ("Price index retrieval manifest", csv(PROJ / "Datasets" / "Raw Macro" / "Revision" / "Price Index Retrieval Manifest.csv")),
]

# ------------------------------------------------------------------ 9. figure data
S["Figure Data"] = [
    ("Figure 2: category prevalence by fiscal year", csv(RES / "Disclosure Trends.csv").rename(columns={"Unnamed: 0": "Fiscal year"})),
    ("Figure 3: breadth dispersion (first version)", csv(RES / "Breadth Dispersion.csv").rename(columns={"Unnamed: 0": "Fiscal year"})),
    ("Figure 4: mean breadth by SIC division", csv(RES / "Breadth by Division.csv").rename(columns={"Unnamed: 0": "Division"})),
    ("Figure 5: estimates plotted (principal models and response variants are in Hypothesis Tests)",
     csv(HERE / "Principal Models Full Inference.csv")[lambda d: d["model"].isin(["H1", "H2", "H3"])]
     [["episode", "outcome", "model", "term", "coef", "ci_low_t", "ci_high_t", "p", "n"]]),
]

# ------------------------------------------------------------------ write
GREY, LIGHT, RULE = "595959", "F2F2F2", "BFBFBF"
FONT = "Calibri"
wb = Workbook()
wb.remove(wb.active)
used = set()
for sname, blocks in S.items():
    ws = wb.create_sheet(sname)
    ws.sheet_view.showGridLines = False
    row = 1
    widths = {}
    for title, df in blocks:
        if df is None or df.empty:
            continue
        df = tidy(df)
        ws.cell(row=row, column=1, value=title).font = Font(name=FONT, size=11, bold=True, color="262626")
        row += 1
        for j, c in enumerate(df.columns, 1):
            cell = ws.cell(row=row, column=j, value=c)
            cell.font = Font(name=FONT, size=10, bold=True, color="262626")
            cell.fill = PatternFill("solid", fgColor=LIGHT)
            cell.border = Border(bottom=Side(style="thin", color=RULE))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[j] = max(widths.get(j, 8), min(len(str(c)) + 2, 40))
        for i, rec_ in enumerate(df.itertuples(index=False), 1):
            for j, v in enumerate(rec_, 1):
                if isinstance(v, float) and pd.isna(v):
                    v = None
                cell = ws.cell(row=row + i, column=j, value=v)
                cell.font = Font(name=FONT, size=10, color="262626")
                cell.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 60)
                if isinstance(v, float):
                    cell.number_format = "0.0000" if abs(v) < 10 else "#,##0.00"
                if v is not None:
                    widths[j] = max(widths.get(j, 8), min(len(str(v)) + 2, 60))
        ref = f"A{row}:{get_column_letter(len(df.columns))}{row + len(df)}"
        base = re.sub(r"[^A-Za-z]", "", sname + title.title())[:40] or "Table"
        name, k = base, 1
        while name in used:
            k += 1
            name = f"{base}{chr(64 + k)}"
        used.add(name)
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False, showColumnStripes=False)
        ws.add_table(t)
        row += len(df) + 3
    for j, w in widths.items():
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
wb.save(OUT)
print("written", OUT, [ws.title for ws in wb.worksheets])
